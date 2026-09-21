from __future__ import annotations

import io
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XlImage

from infrastructure.spreadsheet.instruction_sheet import (
    MissingProductImageError,
    InstructionSheet,
    TEMPLATE_PATH,
    count_embedded_images,
    find_rows_without_image,
    is_usable_product_image,
    template_image_count,
)

PNG_1PX = bytes.fromhex(
    "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c489"
    "0000000d4944415478da63fcffff3f0300050001aeb9c2ff0000000049454e44ae426082"
)
TRANSPARENT_GIF = bytes.fromhex(
    "47494638396101000100800000ffffff00000021f90401000000002c000000000100010000020144003b"
)


def _row(row_number: int, asin: str) -> dict[str, object]:
    return {
        "fnsku": "X001",
        "asin": asin,
        "quantity": "10",
        "remarks": "袋詰め",
        "order_number": "PO-1",
        "row_numbers": [row_number],
    }


class TestIsUsableProductImage:
    def test_十分な大きさのPNGは使える(self) -> None:
        assert is_usable_product_image(PNG_1PX + b"\x00" * 4096) is True

    def test_JPEGも使える(self) -> None:
        assert is_usable_product_image(b"\xff\xd8\xff\xe0" + b"\x00" * 4096) is True

    def test_Noneは使えない(self) -> None:
        assert is_usable_product_image(None) is False

    def test_空のバイト列は使えない(self) -> None:
        assert is_usable_product_image(b"") is False

    def test_透明GIFのプレースホルダは使えない(self) -> None:
        assert is_usable_product_image(TRANSPARENT_GIF) is False

    def test_HTMLのエラーページは使えない(self) -> None:
        assert is_usable_product_image(b"<!DOCTYPE html><html>" + b"x" * 4096) is False

    def test_画像形式でも極端に小さければ使えない(self) -> None:
        assert is_usable_product_image(PNG_1PX) is False


class TestFindRowsWithoutImage:
    def test_全行に画像があれば空を返す(self) -> None:
        rows = [_row(10, "B001"), _row(11, "B002")]
        images = {"B001": b"one", "B002": b"two"}
        assert find_rows_without_image(rows, images) == []

    def test_画像の無い行を仕入管理の行番号とASINで返す(self) -> None:
        rows = [_row(10, "B001"), _row(11, "B002")]
        images = {"B001": b"one"}
        assert find_rows_without_image(rows, images) == [(11, "B002")]

    def test_ASINが空の行も欠落として返す(self) -> None:
        rows = [_row(12, "")]
        assert find_rows_without_image(rows, {}) == [(12, "")]


class TestCountEmbeddedImages:
    def test_テンプレートの埋め込み枚数を数える(self) -> None:
        assert count_embedded_images(TEMPLATE_PATH) == template_image_count()

    def test_追加した画像の分だけ増える(self, tmp_path: Path) -> None:
        workbook = load_workbook(str(TEMPLATE_PATH))
        worksheet = workbook.active
        worksheet.add_image(XlImage(io.BytesIO(PNG_1PX)), "A8")
        saved = tmp_path / "out.xlsx"
        workbook.save(str(saved))
        assert count_embedded_images(saved) == template_image_count() + 1


class TestCreateStopsWithoutImages:
    def _sheet(self, tmp_path: Path, payloads: dict[str, bytes | None]) -> InstructionSheet:
        sheet = InstructionSheet(save_dir=tmp_path, keepa_api_key="")
        sheet._resolve_image_url = lambda asin: f"http://img/{asin}"  # type: ignore[method-assign]
        sheet._fetch_image_bytes = lambda url: payloads.get(url)  # type: ignore[method-assign]
        return sheet

    def _data(self, *asins: str) -> list[object]:
        class Row(dict):
            def __init__(self, row_number: int, asin: str) -> None:
                super().__init__(
                    {
                        "FNSKU": f"X00{row_number}",
                        "SKU": f"SKU-{row_number}",
                        "ASIN": asin,
                        "購入数": "10",
                        "備考": "袋詰め",
                        "注文番号": "PO-1",
                        "納品分類": "ノーマル",
                    }
                )
                self.row_number = row_number

        return [Row(10 + i, asin) for i, asin in enumerate(asins)]

    def test_画像が取れない商品があれば指示書を作らない(self, tmp_path: Path) -> None:
        sheet = self._sheet(tmp_path, {})
        with pytest.raises(MissingProductImageError, match="B001"):
            sheet.create(self._data("B001"))
        assert list(tmp_path.glob("*.xlsx")) == []

    def test_プレースホルダ画像しか取れない商品があれば止める(self, tmp_path: Path) -> None:
        sheet = self._sheet(tmp_path, {"http://img/B001": TRANSPARENT_GIF})
        with pytest.raises(MissingProductImageError, match="B001"):
            sheet.create(self._data("B001"))

    def test_全商品の写真が揃っていれば保存する(self, tmp_path: Path) -> None:
        payload = PNG_1PX + b"\x00" * 4096
        sheet = self._sheet(tmp_path, {"http://img/B001": payload, "http://img/B002": payload})
        path = sheet.create(self._data("B001", "B002"))
        assert path.exists()
        assert count_embedded_images(path) == template_image_count() + 2

    def test_保存した指示書に写真が置かれていなければ止める(self, tmp_path: Path) -> None:
        payload = PNG_1PX + b"\x00" * 4096
        sheet = self._sheet(tmp_path, {"http://img/B001": payload})
        sheet._write_row_data = lambda ws, rows, images: None  # type: ignore[method-assign]
        with pytest.raises(MissingProductImageError, match="写真"):
            sheet.create(self._data("B001"))
        assert list(tmp_path.glob("*.xlsx")) == []


class TestWriteRowDataWithoutImage:
    def test_画像が無い行は画像だけ省いて書く(self) -> None:
        from openpyxl import Workbook

        sheet = InstructionSheet(save_dir=Path("/tmp"), keepa_api_key="")
        worksheet = Workbook().active
        sheet._write_row_data(worksheet, [_row(10, "B001")], images={})
        assert worksheet.cell(row=8, column=3).value == "B001"
        assert worksheet.cell(row=8, column=4).value == 10
