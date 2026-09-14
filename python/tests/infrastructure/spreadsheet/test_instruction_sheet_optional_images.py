from __future__ import annotations

from pathlib import Path

import pytest

from infrastructure.spreadsheet.instruction_sheet import InstructionSheet


def _rows(*asins: str) -> list[dict[str, str]]:
    return [
        {"fnsku": f"X{i}", "asin": a, "quantity": "10", "remarks": "", "order_number": ""}
        for i, a in enumerate(asins)
    ]


def _sheet(urls: dict[str, str | None], payloads: dict[str, bytes | None]) -> InstructionSheet:
    s = InstructionSheet(save_dir=Path("/tmp"), keepa_api_key="")
    s._resolve_image_url = lambda asin: urls.get(asin)  # type: ignore[method-assign]
    s._fetch_image_bytes = lambda url: payloads.get(url)  # type: ignore[method-assign]
    return s


class TestRequireImages:
    def test_既定では画像が無ければ止める(self) -> None:
        sheet = _sheet({"B001": None}, {})
        with pytest.raises(RuntimeError, match="B001"):
            sheet._collect_images(_rows("B001"))

    def test_必須でなければ画像なしでも通す(self) -> None:
        sheet = _sheet({"B001": None}, {})
        assert sheet._collect_images(_rows("B001"), require_images=False) == {}

    def test_必須でなくても取れた分は返す(self) -> None:
        sheet = _sheet(
            {"B001": "http://img/1.jpg", "B002": None}, {"http://img/1.jpg": b"one"}
        )
        got = sheet._collect_images(_rows("B001", "B002"), require_images=False)
        assert got == {"B001": b"one"}

    def test_必須でなければASIN空でも止めない(self) -> None:
        sheet = _sheet({}, {})
        assert sheet._collect_images(_rows(""), require_images=False) == {}


class TestWriteRowDataWithoutImage:
    def test_画像が無い行は画像だけ省いて書く(self) -> None:
        from openpyxl import Workbook

        sheet = InstructionSheet(save_dir=Path("/tmp"), keepa_api_key="")
        ws = Workbook().active
        rows = _rows("B001")
        sheet._write_row_data(ws, rows, images={})  # 画像なし
        assert ws.cell(row=8, column=3).value == "B001"
        assert ws.cell(row=8, column=4).value == 10
